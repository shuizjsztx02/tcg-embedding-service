#!/usr/bin/env python3
"""Batch-test the deployed ``POST /v2/recognize`` endpoint with V3 images and OCR text."""

from __future__ import annotations

import argparse
import csv
import json
import math
import mimetypes
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable

import openpyxl
import requests


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASE_URL = "http://172.31.12.82:8003"
DEFAULT_IMAGES_DIR = PROJECT_ROOT / "test-images-v3"
DEFAULT_OCR_XLSX = PROJECT_ROOT / "OCR识别结果.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "test-results-recognize-v3"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}


@dataclass(frozen=True)
class RecognizeCase:
    filename: str
    image_path: Path
    text: str
    ocr_elapsed_ms: float | None


def read_ocr_rows(xlsx_path: Path) -> list[tuple[Any, Any, Any]]:
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if tuple(header[:3]) != ("图片名", "识别文字", "耗时"):
            raise ValueError(f"unexpected OCR workbook header: {header[:3]!r}")
        return list(sheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()


def build_cases(
    ocr_rows: Iterable[tuple[Any, Any, Any]], image_dir: Path
) -> list[RecognizeCase]:
    image_paths = {
        path.name: path
        for path in image_dir.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    }
    rows_by_filename: dict[str, tuple[str, float | None]] = {}
    for row in ocr_rows:
        filename = str(row[0] or "").strip()
        text = str(row[1] or "")
        if not filename:
            raise ValueError("blank OCR filename")
        if not text.strip():
            raise ValueError(f"blank OCR text: {filename}")
        if filename in rows_by_filename:
            raise ValueError(f"duplicate OCR filename: {filename}")
        elapsed = row[2] if len(row) > 2 else None
        rows_by_filename[filename] = (text, float(elapsed) if elapsed is not None else None)

    ocr_names = set(rows_by_filename)
    image_names = set(image_paths)
    if ocr_names != image_names:
        missing_images = sorted(ocr_names - image_names)
        missing_ocr = sorted(image_names - ocr_names)
        raise ValueError(
            "mapping mismatch: "
            f"xlsx_without_image={len(missing_images)}, image_without_xlsx={len(missing_ocr)}"
        )

    return [
        RecognizeCase(
            filename=filename,
            image_path=image_paths[filename],
            text=rows_by_filename[filename][0],
            ocr_elapsed_ms=rows_by_filename[filename][1],
        )
        for filename in sorted(rows_by_filename)
    ]


def _response_error(response: Any) -> str:
    body = getattr(response, "text", "")
    return f"HTTP {response.status_code}: {str(body)[:500]}"


def call_recognize(
    base_url: str,
    image_path: Path,
    text: str,
    timeout: float,
    post: Callable[..., Any] = requests.post,
) -> dict[str, Any]:
    """Call the compatibility endpoint once with the image and the supplied OCR text."""
    mime_type = mimetypes.guess_type(image_path.name)[0] or "image/jpeg"
    image_bytes = image_path.read_bytes()
    started = time.perf_counter()
    try:
        response = post(
            f"{base_url.rstrip('/')}/v2/recognize",
            files={"image": (image_path.name, image_bytes, mime_type)},
            data={"text": text},
            timeout=timeout,
        )
    except requests.RequestException as exc:
        return {
            "filename": image_path.name,
            "http_status": None,
            "http_ok": False,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 1),
            "raw_response": None,
            "error": f"request failed: {exc}",
        }
    elapsed_ms = round((time.perf_counter() - started) * 1000, 1)
    record: dict[str, Any] = {
        "filename": image_path.name,
        "http_status": response.status_code,
        "http_ok": bool(getattr(response, "ok", False)),
        "elapsed_ms": elapsed_ms,
        "raw_response": None,
        "error": None,
    }
    if not record["http_ok"]:
        record["error"] = _response_error(response)
        return record

    try:
        payload = response.json()
    except ValueError:
        record["error"] = "invalid JSON response"
        return record

    text_result = payload.get("text") if isinstance(payload, dict) else {}
    monitor = payload.get("monitor") if isinstance(payload, dict) else {}
    text_result = text_result if isinstance(text_result, dict) else {}
    monitor = monitor if isinstance(monitor, dict) else {}
    record.update(
        status_match=monitor.get("status_match", ""),
        strategy=monitor.get("strategy", ""),
        service_latency_ms=monitor.get("latency_ms"),
        product_id=text_result.get("productId"),
        card_name=text_result.get("cardName", ""),
        card_ip=text_result.get("cardIp", ""),
        data_source=monitor.get("dataSource", ""),
        raw_response=payload,
    )
    return record


def process_case(case: RecognizeCase, base_url: str, timeout: float) -> dict[str, Any]:
    try:
        record = call_recognize(base_url, case.image_path, case.text, timeout)
    except (OSError, requests.RequestException) as exc:
        record = {
            "filename": case.filename,
            "http_status": None,
            "http_ok": False,
            "elapsed_ms": 0.0,
            "raw_response": None,
            "error": f"request failed: {exc}",
        }
    record["ocr_elapsed_ms"] = case.ocr_elapsed_ms
    record["ocr_text_chars"] = len(case.text)
    record["tested_at"] = datetime.now().isoformat(timespec="seconds")
    return record


def load_completed_filenames(results_path: Path) -> set[str]:
    if not results_path.is_file():
        return set()
    completed: set[str] = set()
    with results_path.open("r", encoding="utf-8") as results_file:
        for line in results_file:
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue
            filename = record.get("filename")
            if filename:
                completed.add(filename)
    return completed


def append_jsonl(results_path: Path, record: dict[str, Any]) -> None:
    needs_newline = results_path.exists() and results_path.stat().st_size > 0
    if needs_newline:
        with results_path.open("rb") as existing_file:
            existing_file.seek(-1, 2)
            needs_newline = existing_file.read(1) != b"\n"
    with results_path.open("a", encoding="utf-8") as results_file:
        if needs_newline:
            results_file.write("\n")
        results_file.write(json.dumps(record, ensure_ascii=False) + "\n")


def _percentile(values: list[float], percentile: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    position = (len(ordered) - 1) * percentile
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return float(ordered[lower])
    return round(ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower), 1)


def summarize_records(records: Iterable[dict[str, Any]], expected_total: int) -> dict[str, Any]:
    all_records = list(records)
    successful = [record for record in all_records if record.get("http_ok") and not record.get("error")]
    latencies = [float(record["elapsed_ms"]) for record in successful if record.get("elapsed_ms") is not None]
    return {
        "expected_total": expected_total,
        "processed": len(all_records),
        "http_success": len(successful),
        "errors": len(all_records) - len(successful),
        "status_match": dict(sorted(Counter(record.get("status_match", "unknown") for record in successful).items())),
        "strategy": dict(sorted(Counter(record.get("strategy", "unknown") for record in successful).items())),
        "elapsed_ms": {
            "avg": round(sum(latencies) / len(latencies), 1) if latencies else 0.0,
            "p50": _percentile(latencies, 0.50),
            "p95": _percentile(latencies, 0.95),
            "p99": _percentile(latencies, 0.99),
            "max": max(latencies) if latencies else 0.0,
        },
    }


def load_records(results_path: Path) -> list[dict[str, Any]]:
    if not results_path.is_file():
        return []
    records: list[dict[str, Any]] = []
    with results_path.open("r", encoding="utf-8") as results_file:
        for line in results_file:
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    return records


def write_csv(records: Iterable[dict[str, Any]], csv_path: Path) -> None:
    columns = [
        "filename", "http_status", "http_ok", "elapsed_ms", "service_latency_ms",
        "status_match", "strategy", "product_id", "card_name", "card_ip",
        "data_source", "ocr_elapsed_ms", "ocr_text_chars", "error",
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=columns)
        writer.writeheader()
        writer.writerows({column: record.get(column, "") for column in columns} for record in records)


def check_server(base_url: str) -> bool:
    try:
        return requests.get(f"{base_url.rstrip('/')}/v1/health", timeout=10).ok
    except requests.RequestException:
        return False


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Batch-test /v2/recognize with V3 images and OCR text")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help="recognize service base URL")
    parser.add_argument("--images-dir", type=Path, default=DEFAULT_IMAGES_DIR)
    parser.add_argument("--ocr-xlsx", type=Path, default=DEFAULT_OCR_XLSX)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--workers", type=int, default=4, help="number of concurrent requests")
    parser.add_argument("--timeout", type=float, default=180.0, help="per-request timeout in seconds")
    parser.add_argument("--limit", type=int, help="test only the first N sorted images")
    parser.add_argument("--resume", action="store_true", help="skip filenames already present in JSONL output")
    parser.add_argument("--overwrite", action="store_true", help="replace an existing JSONL output")
    parser.add_argument("--skip-health-check", action="store_true")
    args = parser.parse_args()
    if args.workers < 1:
        parser.error("--workers must be at least 1")
    if args.timeout <= 0:
        parser.error("--timeout must be positive")
    if args.limit is not None and args.limit < 1:
        parser.error("--limit must be at least 1")
    if args.resume and args.overwrite:
        parser.error("--resume and --overwrite cannot be used together")
    return args


def main() -> int:
    args = parse_args()
    if not args.images_dir.is_dir():
        raise SystemExit(f"image directory not found: {args.images_dir}")
    if not args.ocr_xlsx.is_file():
        raise SystemExit(f"OCR workbook not found: {args.ocr_xlsx}")
    if not args.skip_health_check and not check_server(args.base_url):
        raise SystemExit(f"service health check failed: {args.base_url}/v1/health")

    cases = build_cases(read_ocr_rows(args.ocr_xlsx), args.images_dir)
    if args.limit is not None:
        cases = cases[:args.limit]
    args.output.mkdir(parents=True, exist_ok=True)
    results_path = args.output / "recognize_results.jsonl"
    if results_path.exists() and not args.resume and not args.overwrite:
        raise SystemExit(f"results already exist: {results_path}; use --resume or --overwrite")
    if args.overwrite:
        results_path.write_text("", encoding="utf-8")

    completed = load_completed_filenames(results_path) if args.resume else set()
    pending = [case for case in cases if case.filename not in completed]
    config = {
        "base_url": args.base_url,
        "endpoint": f"{args.base_url.rstrip('/')}/v2/recognize",
        "images_dir": str(args.images_dir),
        "ocr_xlsx": str(args.ocr_xlsx),
        "workers": args.workers,
        "timeout": args.timeout,
        "selected_images": len(cases),
        "pending_images": len(pending),
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    (args.output / "config.json").write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Testing {len(pending)}/{len(cases)} images against {config['endpoint']}", flush=True)
    started = time.perf_counter()
    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {executor.submit(process_case, case, args.base_url, args.timeout): case.filename for case in pending}
        for index, future in enumerate(as_completed(futures), start=1):
            record = future.result()
            append_jsonl(results_path, record)
            if index % 25 == 0 or index == len(pending):
                print(f"  {index}/{len(pending)} complete", flush=True)

    records = load_records(results_path)
    summary = summarize_records(records, expected_total=len(cases))
    summary["run_duration_seconds"] = round(time.perf_counter() - started, 1)
    (args.output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(records, args.output / "recognize_results.csv")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
